# -*- coding: utf-8 -*-
"""Gera o documento da Entrega Parcial A3 (ABNT) em .docx.
Uso: python build_documento.py
"""
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(AQUI, "img")
SAIDA = os.path.join(AQUI, "Entrega_Parcial_A3_Sandbox_Grafico_Interativo.docx")
RES = json.load(open(os.path.join(AQUI, "resultados.json"), encoding="utf-8"))

FONTE = "Arial"

doc = Document()

# ---------------------------------------------------------------- estilos base
st = doc.styles["Normal"]
st.font.name = FONTE
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)

for nome, caps in (("Heading 1", True), ("Heading 2", False), ("Heading 3", False)):
    h = doc.styles[nome]
    h.font.name = FONTE
    h.font.size = Pt(12)
    h.font.bold = True
    h.font.italic = False
    h.font.color.rgb = None
    h.font.all_caps = caps
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    h.element.rPr.rFonts.set(qn("w:ascii"), FONTE)
    h.element.rPr.rFonts.set(qn("w:hAnsi"), FONTE)
    h.element.rPr.rFonts.set(qn("w:cs"), FONTE)
    for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        k = qn("w:" + att)
        if k in h.element.rPr.rFonts.attrib:
            del h.element.rPr.rFonts.attrib[k]
    h.paragraph_format.line_spacing = 1.5
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(12)
    h.paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.left_margin = Cm(3), Cm(3)
sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)


# ---------------------------------------------------------------- utilitários
def _runs(p, texto, bold=False, italic=False, size=None):
    """Aceita marcação simples: _{sub}, **negrito**."""
    partes = re.split(r"(_\{.*?\}|\*\*.*?\*\*)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("_{"):
            r = p.add_run(parte[2:-1])
            r.font.subscript = True
        elif parte.startswith("**"):
            r = p.add_run(parte[2:-2])
            r.bold = True
        else:
            r = p.add_run(parte)
            r.bold = bold
        r.italic = italic
        r.font.name = FONTE
        if size:
            r.font.size = Pt(size)
    return p


def par(texto, recuo=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=None,
        spacing=1.5, after=0, before=0, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.line_spacing = spacing
    f.space_after = Pt(after)
    f.space_before = Pt(before)
    f.keep_with_next = keep
    if recuo:
        f.first_line_indent = Cm(1.25)
    _runs(p, texto, bold=bold, italic=italic, size=size)
    return p


def h1(t):
    return doc.add_heading(t, level=1)


def h2(t):
    return doc.add_heading(t, level=2)


def citacao_longa(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    _runs(p, texto, size=10)


def equacao(texto, num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _runs(p, texto + "        (" + str(num) + ")")


def legenda(texto):
    p = par(texto, recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, spacing=1.0, before=6, after=4, keep=True)
    return p


def fonte(texto="Fonte: Elaborado pelos autores (2026)."):
    return par(texto, recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, spacing=1.0, after=12)


def figura(arquivo, largura_cm, titulo, fonte_txt="Fonte: Elaborado pelos autores (2026)."):
    legenda(titulo)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(os.path.join(IMG, arquivo), width=Cm(largura_cm))
    fonte(fonte_txt)


def _borda(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado, v in kw.items():
        e = OxmlElement("w:" + lado)
        e.set(qn("w:val"), v)
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tcPr.append(b)


def tabela(cabecalho, linhas, larguras_cm, titulo, fonte_txt="Fonte: Elaborado pelos autores (2026).", size=10):
    legenda(titulo)
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, c in enumerate(cabecalho):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(p, c, bold=True, size=size)
    for linha in linhas:
        row = t.add_row()
        for i, c in enumerate(linha):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 or len(c) < 14 else WD_ALIGN_PARAGRAPH.LEFT
            _runs(p, c, size=size)
    for ri, row in enumerate(t.rows):
        for i, w in enumerate(larguras_cm):
            row.cells[i].width = Cm(w)
            if ri < len(t.rows) - 1:
                for pp in row.cells[i].paragraphs:
                    pp.paragraph_format.keep_with_next = True
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    fonte(fonte_txt)
    return t


def campo(par_, instr):
    r = par_.add_run()
    for tipo, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        elif txt == instr:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = instr
        else:
            e = OxmlElement("w:t")
            e.text = txt
        r._r.append(e)
    r.font.name = FONTE
    r.font.size = Pt(10)


def quebra():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================ CAPA
par("UNICURITIBA", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.0)
par("CURSO: [PREENCHER O NOME DO CURSO]", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.0)
par("COMPUTAÇÃO GRÁFICA E REALIDADE VIRTUAL", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=6)
par("", recuo=False, spacing=1.0, before=70)
par("Gabriel Thiago Casini Keppen (RA 172422353)", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("Felipe Silva Muchinski (RA 172422440)", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("", recuo=False, spacing=1.0, before=90)
par("SANDBOX GRÁFICO INTERATIVO PARA VISUALIZAÇÃO DE ILUMINAÇÃO E CONCEITOS MATEMÁTICOS EM COMPUTAÇÃO "
    "GRÁFICA UTILIZANDO UNREAL ENGINE", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.5)
par("Entrega parcial do Projeto A3", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, before=12)
par("", recuo=False, spacing=1.0, before=110)
par("Curitiba", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("2026", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
quebra()

# ================================================================ FOLHA DE ROSTO
par("Gabriel Thiago Casini Keppen", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("Felipe Silva Muchinski", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("", recuo=False, spacing=1.0, before=110)
par("SANDBOX GRÁFICO INTERATIVO PARA VISUALIZAÇÃO DE ILUMINAÇÃO E CONCEITOS MATEMÁTICOS EM COMPUTAÇÃO "
    "GRÁFICA UTILIZANDO UNREAL ENGINE", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.5)
par("", recuo=False, spacing=1.0, before=60)
p = par("Documento referente à entrega parcial do Projeto A3 da unidade curricular Computação Gráfica e "
        "Realidade Virtual, apresentado como requisito de avaliação.", recuo=False, spacing=1.0, size=10)
p.paragraph_format.left_indent = Cm(8)
p = par("Professor: Rubem Matimoto Koide", recuo=False, spacing=1.0, size=10, before=10)
p.paragraph_format.left_indent = Cm(8)
par("", recuo=False, spacing=1.0, before=120)
par("Curitiba", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
par("2026", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
quebra()

# ================================================================ SUMÁRIO
par("SUMÁRIO", recuo=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=12)
p = doc.add_paragraph()
r = p.add_run()
for tipo, txt in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                  ("txt", "Clique com o botão direito e escolha “Atualizar campo” para gerar o sumário."), ("end", None)):
    if tipo in ("begin", "separate", "end"):
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), tipo)
    elif tipo == "instr":
        e = OxmlElement("w:instrText")
        e.set(qn("xml:space"), "preserve")
        e.text = txt
    else:
        e = OxmlElement("w:t")
        e.text = txt
    r._r.append(e)

# nova seção: a numeração aparece a partir da introdução
nova = doc.add_section(WD_SECTION.NEW_PAGE)
hdr = nova.header
hdr.is_linked_to_previous = False
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
campo(hp, "PAGE")

# ================================================================ 1 INTRODUÇÃO
h1("1 Introdução")
h2("1.1 Contextualização")
par("A computação gráfica utiliza, de forma contínua, conceitos matemáticos para representar, transformar e iluminar "
    "objetos em ambientes tridimensionais. Operações com vetores, matrizes, transformações geométricas e produtos "
    "escalares estão presentes em cada quadro renderizado (HUGHES et al., 2013; AKENINE-MÖLLER; HAINES; HOFFMAN, 2018). "
    "Nos motores gráficos atuais, no entanto, boa parte desse cálculo ocorre de maneira automática e fica oculta de "
    "quem desenvolve ou estuda a área, o que dificulta relacionar a teoria vista em sala com o resultado visual obtido.")
h2("1.2 Problema e proposta")
par("Como tornar perceptível, em tempo real, a relação entre a matemática da iluminação e a aparência final de uma "
    "superfície em um motor gráfico? Para responder a essa pergunta, o projeto propõe um **sandbox gráfico interativo** "
    "desenvolvido na Unreal Engine, no qual o usuário manipula uma fonte de luz e observa, sobre a própria cena, o vetor "
    "normal da superfície, a direção da luz, o produto escalar entre eles, o ângulo de incidência e a intensidade "
    "difusa resultante.")
h2("1.3 Objetivos")
par("**Objetivo geral:** desenvolver um ambiente 3D interativo que permita visualizar, em tempo real, os cálculos "
    "vetoriais envolvidos na iluminação difusa de uma superfície.")
par("**Objetivos específicos:**", recuo=False, after=0)
for t in ("construir uma cena 3D na Unreal Engine com superfície, piso e fonte de luz manipulável;",
          "implementar em Blueprint o cálculo do produto escalar N · L, do ângulo de incidência e da intensidade "
          "difusa;",
          "representar graficamente os vetores normal e de direção da luz diretamente na cena;",
          "validar os valores calculados pelo Blueprint contra o cálculo teórico;",
          "projetar a estrutura de modo modular, permitindo incluir futuramente reflexão especular, transformações, "
          "matrizes, materiais e sombras."):
    p = par("• " + t, recuo=False)
    p.paragraph_format.left_indent = Cm(1.25)
h2("1.4 Justificativa")
par("O projeto tem valor didático: ao expor valores normalmente invisíveis, aproxima o formalismo matemático da "
    "experiência visual e funciona como ferramenta de apoio ao estudo da disciplina. Também se alinha às aplicações "
    "de computação gráfica e realidade virtual discutidas em aula, como visualização técnica, simulação e "
    "treinamento (KOIDE, 2026).")
h2("1.5 Metodologia")
par("O desenvolvimento segue abordagem incremental, partindo de um produto mínimo viável (cena, luz manipulável, "
    "visualização dos vetores e cálculo da iluminação difusa) e incorporando módulos conforme o tempo permite. O "
    "trabalho é versionado em repositório Git hospedado no GitHub. Na criação do Blueprint e da cena foi utilizado "
    "um assistente de inteligência artificial (Claude Code), conectado ao editor por meio do Model Context Protocol, "
    "como apoio à automação de tarefas do editor; o projeto, as decisões de modelagem e a validação dos resultados "
    "foram conduzidos pela equipe.")

# ================================================================ 2 REVISÃO
h1("2 Revisão bibliográfica")
h2("2.1 Iluminação difusa e o modelo de Lambert")
par("Superfícies foscas espalham a luz recebida de maneira aproximadamente igual em todas as direções. Essa "
    "propriedade é descrita pela lei do cosseno de Lambert, segundo a qual a luz refletida por um ponto da superfície "
    "é proporcional ao cosseno do ângulo entre a normal da superfície e a direção da luz incidente (LAMBERT, 1760). "
    "Na computação gráfica, esse é o modelo de reflexão difusa mais elementar e constitui o termo difuso de modelos "
    "de iluminação local mais completos (HUGHES et al., 2013). Em sua forma usual:")
equacao("I_{d} = I_{L} · k_{d} · máx(0, N · L)", 1)
par("em que I_{d} é a intensidade difusa resultante, I_{L} a intensidade da luz, k_{d} o coeficiente de "
    "reflexão difusa do material, N o vetor normal unitário da superfície e L o vetor unitário que aponta do ponto "
    "da superfície para a fonte de luz. A função máx impede contribuição negativa quando a luz incide pela face "
    "oposta da superfície.", recuo=False)
h2("2.2 Vetores, normal e produto escalar")
par("O produto escalar entre dois vetores a e b é definido como a · b = |a| |b| cos θ, em que θ é o ângulo entre "
    "eles. Para vetores unitários, o módulo é igual a 1 e o produto escalar coincide com o cosseno do ângulo "
    "(AKENINE-MÖLLER; HAINES; HOFFMAN, 2018). Assim, o ângulo de incidência da luz pode ser recuperado por:")
equacao("θ = arccos(N · L)", 2)
par("Para uma luz pontual na posição P_{L} e um ponto da superfície na posição P, o vetor de direção da luz é "
    "obtido pela normalização da diferença entre as posições:", recuo=False)
equacao("L = (P_{L} − P) / |P_{L} − P|", 3)
h2("2.3 Unreal Engine e Blueprints")
par("A Unreal Engine, desenvolvida pela Epic Games, é um motor gráfico de tempo real amplamente utilizado em jogos, "
    "visualização arquitetônica, simulação e produção virtual. Seu sistema de programação visual, o Blueprint, "
    "permite construir lógica de interação por meio de nós conectados, sem escrever código tradicional, com suporte "
    "a C++ quando necessário (EPIC GAMES, 2026). Esse recurso foi escolhido por permitir iterar rapidamente sobre "
    "o cálculo e por tornar a lógica legível para fins didáticos, conforme as ferramentas apresentadas em aula "
    "(KOIDE, 2026).")
h2("2.4 Ferramentas de apoio ao desenvolvimento")
par("**Git, GitHub e Git LFS.** O controle de versão é feito com Git (CHACON; STRAUB, 2014), com repositório remoto "
    "no GitHub. Como projetos Unreal possuem arquivos binários grandes (.uasset e .umap), é utilizado o Git LFS, "
    "que armazena esses arquivos fora do histórico regular do repositório.")
par("**Model Context Protocol (MCP).** O MCP é um protocolo aberto que padroniza a conexão entre assistentes de "
    "inteligência artificial e ferramentas externas (MODEL CONTEXT PROTOCOL, 2026). A Unreal Engine 5.8 inclui, "
    "como recurso experimental, o plugin ModelContextProtocol, que faz o editor atuar como servidor MCP e expõe "
    "ferramentas para manipular atores, níveis, assets e grafos de Blueprint. No projeto, esse plugin foi ativado "
    "(juntamente com os toolsets EditorToolset e UMGToolSet) para automatizar a criação da cena e do Blueprint.")

# ================================================================ 3 DESENVOLVIMENTO PARCIAL
h1("3 Desenvolvimento parcial")
h2("3.1 Ambiente de desenvolvimento e organização do projeto")
par("O projeto foi criado na Unreal Engine 5.8 com o nome LightSandbox, a partir de um modelo em branco baseado em "
    "Blueprint. O código-fonte e os assets são versionados no repositório privado da equipe no GitHub "
    "(github.com/Gabrielkeppen/Light_SandBox), com Git LFS habilitado para os assets binários. O Quadro 1 resume o "
    "ambiente.")
tabela(["Item", "Descrição"],
       [["Motor gráfico", "Unreal Engine 5.8"],
        ["Linguagem / lógica", "Blueprint (programação visual)"],
        ["Plugins ativados", "ModelContextProtocol, EditorToolset, UMGToolSet, MCPClientToolset"],
        ["Controle de versão", "Git, GitHub (repositório privado) e Git LFS"],
        ["Sistema operacional", "Windows 10"]],
       [5, 9.5], "Quadro 1 – Ambiente de desenvolvimento")
h2("3.2 Cena de testes")
par("A cena de testes (nível SandBoxMap) é composta por três elementos, descritos no Quadro 2: um piso, uma luz "
    "pontual e o ator analisador, que representa a superfície cujo vetor normal e iluminação são estudados.")
tabela(["Elemento", "Tipo", "Posição (cm)", "Observações"],
       [["Chao", "Static Mesh (plano)", "(0, 0, 0)", "Escala 20 × 20; referência visual"],
        ["LuzPontual", "Point Light", "(350, −250, 400)", "Fonte de luz manipulável"],
        ["Analisador", "BP_LightAnalyzer", "(0, 0, 150)", "Inclinado 20° (pitch); referencia a LuzPontual"]],
       [3, 3.6, 3.4, 4.5], "Quadro 2 – Elementos da cena")
h2("3.3 Blueprint BP_LightAnalyzer")
par("O Blueprint BP_LightAnalyzer é um ator que contém uma malha em forma de cubo achatado (4 × 4 × 0,1 unidades "
    "de escala) usada como superfície. As variáveis do Blueprint estão no Quadro 3.")
tabela(["Variável", "Tipo", "Função"],
       [["LightActor", "Referência a ator", "Luz analisada; editável por instância na cena"],
        ["NdotL", "Ponto flutuante", "Produto escalar N · L"],
        ["AngleDeg", "Ponto flutuante", "Ângulo de incidência θ, em graus"],
        ["DiffusePct", "Ponto flutuante", "Intensidade difusa, em porcentagem"]],
       [3.2, 3.8, 7.5], "Quadro 3 – Variáveis do Blueprint")
par("A cada quadro (evento Tick), e desde que a referência à luz seja válida, o Blueprint executa os passos abaixo, "
    "que implementam diretamente as equações 1 a 3:")
passos = ["obtém a posição do ator (P) e a posição da luz (P_{L});",
          "calcula L = normalizar(P_{L} − P);",
          "obtém a normal N como o vetor “para cima” (Up Vector) do ator;",
          "calcula N · L com o nó Dot Product;",
          "calcula θ = arccos(N · L), em graus, com o nó Acos (Degrees);",
          "calcula a intensidade difusa como máx(0, N · L) × 100;",
          "desenha, com Draw Debug Arrow, a seta verde da normal e a seta amarela da direção da luz;",
          "exibe os valores na tela e no log (Print String)."]
for i, t in enumerate(passos, 1):
    p = par(f"{i}) " + t, recuo=False)
    p.paragraph_format.left_indent = Cm(1.25)
par("A Figura 1 mostra o ambiente com a simulação em execução e o log com os valores calculados. A Figura 2 "
    "apresenta a superfície inclinada com as setas desenhadas: em verde a normal N e, em amarelo, a direção da luz L.",
    before=6)
figura("01_editor_pie_n_dot_l.png", 14.5, "Figura 1 – Editor da Unreal Engine com simulação em execução e log do cálculo")
figura("02_pie_vista_analisador.png", 14.5, "Figura 2 – Superfície analisada com a normal (verde) e a direção da luz (amarelo)")
h2("3.4 Resultados parciais e validação")
par("Para validar a implementação, a luz foi posicionada em cinco pontos distintos e os valores exibidos pelo "
    "Blueprint foram comparados com o cálculo teórico das equações 1 a 3, feito de forma independente. Com o "
    "analisador em (0, 0, 150) e inclinado 20°, a normal vale N = (−0,342; 0; 0,940). O Quadro 4 apresenta os "
    "resultados.")
linhas = []
for c in RES:
    pos = "(%d, %d, %d)" % tuple(c["luz"])
    linhas.append([c["caso"], pos,
                   f"{c['esperado_ndotl']:.4f}".replace(".", ","), f"{c['ndotl']:.4f}".replace(".", ","),
                   f"{c['esperado_angulo']:.1f}".replace(".", ",") + "°", f"{c['angulo']:.1f}".replace(".", ",") + "°",
                   f"{c['difusa']:.1f}".replace(".", ",") + "%"])
tabela(["Caso", "Posição da luz (cm)", "N·L teórico", "N·L medido", "θ teórico", "θ medido", "Difusa medida"],
       linhas, [1.4, 3.6, 2.1, 2.1, 1.9, 1.9, 2.2], "Quadro 4 – Comparação entre valores teóricos e medidos", size=9)
par("Em todos os casos os valores medidos coincidiram com os teóricos até a sexta casa decimal, o que valida a "
    "implementação do cálculo. O caso C ilustra o comportamento do termo máx(0, N · L): a luz incide pela face "
    "oposta (N · L = −0,342; θ = 110°) e, por isso, a intensidade difusa é 0%. A Figura 3 resume a comparação "
    "para a intensidade difusa.")
figura("03_grafico_resultados.png", 14, "Figura 3 – Intensidade difusa calculada e medida em cada caso")
h2("3.5 Limitações e próximas etapas")
par("O que está implementado cobre o produto mínimo viável do módulo de iluminação difusa. Permanecem como "
    "próximas etapas: (a) interface gráfica (UMG) com os valores em texto sobre a tela, em substituição ao texto de "
    "depuração; (b) controles para o usuário mover a luz durante a execução; (c) inclusão do coeficiente do material "
    "(k_{d}) e da intensidade da luz (I_{L}) no cálculo da equação 1; (d) suporte a luz direcional, além da "
    "pontual; e (e) novos módulos, como reflexão especular, transformações e matrizes, sombras e propriedades de "
    "materiais. Como limitação atual, a normal é obtida do vetor “para cima” do ator, o que é adequado a superfícies "
    "planas, mas não considera malhas com normais variáveis por vértice.")

# ================================================================ REFERÊNCIAS
h1("Referências")
refs = [
    "AKENINE-MÖLLER, Tomas; HAINES, Eric; HOFFMAN, Naty. **Real-time rendering**. 4. ed. Boca Raton: CRC Press, 2018.",
    "CHACON, Scott; STRAUB, Ben. **Pro Git**. 2. ed. Berkeley: Apress, 2014. Disponível em: https://git-scm.com/book. "
    "Acesso em: 12 out. 2026.",
    "EPIC GAMES. **Unreal Engine documentation**. 2026. Disponível em: "
    "https://dev.epicgames.com/documentation/unreal-engine. Acesso em: 12 out. 2026.",
    "HUGHES, John F. et al. **Computer graphics**: principles and practice. 3. ed. Upper Saddle River: "
    "Addison-Wesley, 2013.",
    "KOIDE, Rubem Matimoto. **Computação gráfica e realidade virtual**: aula 1. Curitiba: UniCuritiba, 2026. "
    "Material de aula.",
    "LAMBERT, Johann Heinrich. **Photometria sive de mensura et gradibus luminis, colorum et umbrae**. Augsburg: "
    "Eberhard Klett, 1760.",
    "MODEL CONTEXT PROTOCOL. **Specification**. 2026. Disponível em: https://modelcontextprotocol.io. "
    "Acesso em: 12 out. 2026.",
]
for r in refs:
    par(r, recuo=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0, after=12)

doc.save(SAIDA)
print("gerado:", SAIDA)
